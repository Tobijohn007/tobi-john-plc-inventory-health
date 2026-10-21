"""Writes the ranked action list (colour-coded Excel) from the latest pipeline results."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import config as C

FILLS = {"Excess": "FFE699", "Short": "F4A6A6", "In range": "B7E1B0"}


def export(path=None):
    a = pd.read_csv(C.OUT_DIR / "adequacy_business_units.csv", parse_dates=["month"])
    l = a[a["month"] == a["month"].max()].copy()
    l["total_gap"] = l["excess_value"] + l["shortfall_value"]
    l = l.sort_values("total_gap", ascending=False)
    b = 1e9
    cols = ["Rank", "Business unit", "Vertical", "Status", "Stock now", "Range low", "Range high",
            "Excess above range", "Shortfall below range", "Holding cost of excess per month"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Action list"
    ws.append(cols)
    for i, r in enumerate(l.itertuples(), 1):
        ws.append([i, r.business_unit, r.vertical, r.status, r.closing_value / b, r.range_low / b, r.range_high / b,
                   r.excess_value / b, r.shortfall_value / b, r.excess_holding_cost_month / 1e6])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for row in range(2, ws.max_row + 1):
        fill = PatternFill("solid", fgColor=FILLS[ws.cell(row, 4).value])
        for col in range(1, len(cols) + 1):
            cell = ws.cell(row, col)
            cell.fill = fill
            if 5 <= col <= 9:
                cell.number_format = '"₦"#,##0.00"bn"'
            if col == 10:
                cell.number_format = '"₦"#,##0.0"m"'
        ws.cell(row, 4).font = Font(bold=True)
    ws.row_dimensions[1].height = 45
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, min(26, max(len(str(c.value or "")) for c in col) + 3))
    ws.freeze_panes = "A2"
    path = path or C.OUT_DIR / "ranked_action_list.xlsx"
    wb.save(path)
    return path, l
